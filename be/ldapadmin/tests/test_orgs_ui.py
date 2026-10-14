''' test organisations UI '''

import re
import unittest
import logging
from io import BytesIO, StringIO
from openpyxl import load_workbook
from lxml.html.soupparser import fromstring
from unittest.mock import MagicMock as Mock, patch
from be import ldapadmin
from be.ldapadmin.orgs_editor import OrganisationsEditor, CommonTemplateLogic
from be.ldapadmin.orgs_editor import validate_org_info, VALIDATION_ERRORS
from be.ldapadmin.ui_common import TemplateRenderer

org_info_fixture = {
    'name': u"Ye olde bridge club",
    'phone': u"+45 555 2222",
    'fax': u"+45 555 9999",
    'url': u"http://bridge.example.com/",
    'street': u"Card games road",
    'po_box': u"123456",
    'postal_code': u"DK 456789",
    'country': "dk",
    'locality': u"K\xf8benhavn",
    'email': None,
}



def plaintext(element):
    return re.sub(r'\s\s+', ' ', element.text_content()).strip()

def parse_html(html):
    return fromstring(html)


class StubbedOrganisationsEditor(OrganisationsEditor):
    def __init__(self, id):
        super(StubbedOrganisationsEditor, self).__init__()
        self._render_template = TemplateRenderer(CommonTemplateLogic)
        self._render_template.wrap = lambda html: "<html>%s</html>" % html
        # the changelog is a browser view, which needs a real request
        self._render_template.browserview = lambda context, name: lambda: ''

    def absolute_url(self):
        return "URL"


def mock_request():
    request = Mock()
    request.SESSION = {}
    return request


validation_errors_fixture = {
    'id': [u"invalid ID"],
    'phone': [u"invalid PHONE"],
    'fax': [u"invalid FAX"],
    'postal_code': [u"invalid POSTAL CODE"],
}


def session_messages(request):
    return request.SESSION.get('be.ldapadmin.orgs_editor.messages')


class OrganisationsUITest(unittest.TestCase):
    def setUp(self):
        countries = patch.dict(
            'be.ldapadmin.countries.COUNTRIES',
            {'dk': {'code': 'dk', 'name': 'Denmark', 'pub_code': 'DK'}})
        countries.start()
        self.addCleanup(countries.stop)
        self.ui = StubbedOrganisationsEditor('organisations')
        self.mock_agent = Mock()
        self.mock_agent.all_organisations.return_value = {}
        self.ui._get_ldap_agent = Mock(return_value=self.mock_agent)
        self.request = mock_request()
        user = self.request.AUTHENTICATED_USER
        user.getRoles.return_value = ['Authenticated']
        self.ui.REQUEST = self.request

        self.stream = StringIO()
        self.handler = logging.StreamHandler(self.stream)
        self.log = logging.getLogger('orgs_editor')
        self.log.setLevel(logging.INFO)
        for handler in self.log.handlers:
            self.log.removeHandler(handler)
        self.log.addHandler(self.handler)

    def tearDown(self):
        self.log = logging.getLogger('orgs_editor')
        self.log.removeHandler(self.handler)
        self.handler.close()

    def test_create_org_form(self):
        page = parse_html(self.ui.create_organisation_html(self.request))

        # txt = lambda xp: page.xpath(xp)[0].text.strip()
        exists = lambda xp: len(page.xpath(xp)) > 0
        self.assertTrue(exists('//form//input[@name="id:utf8:string"]'))
        self.assertTrue(exists('//form//input[@name="email:utf8:string"]'))
        self.assertTrue(exists('//form//input[@name="name:utf8:string"]'))
        self.assertTrue(exists('//form//input[@name="url:utf8:string"]'))
        self.assertTrue(exists('//form//input[@name="phone:utf8:string"]'))
        self.assertTrue(exists('//form//input[@name="fax:utf8:string"]'))
        self.assertTrue(exists('//form//input[@name="street:utf8:string"]'))
        self.assertTrue(exists('//form//input[@name="po_box:utf8:string"]'))
        self.assertTrue(exists('//form//input'
                               '[@name="postal_code:utf8:string"]'))
        self.assertTrue(exists('//form//input[@name="locality:utf8:string"]'))
        self.assertTrue(exists('//form//select[@name="country:utf8:string"]'))

    @patch('be.ldapadmin.orgs_editor.logged_in_user')
    def test_create_org_submit(self, logged_user):
        logged_user.return_value = "John Doe"
        self.request.form = dict(org_info_fixture)
        self.request.form['id'] = 'dk_bridge_club'

        self.ui.create_organisation(self.request)

        self.mock_agent.create_org.assert_called_once_with('dk_bridge_club',
                                                           org_info_fixture)
        self.request.RESPONSE.redirect.assert_called_with(
            'URL/organisation?id=dk_bridge_club')

        self.request.form = {'id': 'dk_bridge_club'}
        page = parse_html(self.ui.organisation(self.request))
        self.assertEqual(page.xpath('//div[@class="system-msg"]')[0].text,
                         'Organisation "dk_bridge_club" created successfully.')
        logmsg = "John Doe CREATED ORGANISATION dk_bridge_club\n"
        self.assertEqual(self.stream.getvalue(), logmsg)

    def _verify_org_form_submit_error(self, page, org_info, errors):
        err_msg = page.xpath('//div[@class="error-msg"]')
        self.assertEqual(set(err_div.text for err_div in err_msg), errors)
        self._verify_org_form_values(page.xpath('//form')[0], org_info)

    def _verify_org_form_values(self, form, org_info):
        for name, value in org_info.items():
            if value is None or name == 'id':
                continue
            if name == 'country':
                xp = './/select[@name="country:utf8:string"]/option[@selected]'
            else:
                xp = './/input[@name="%s:utf8:string"]' % name
            self.assertEqual(form.xpath(xp)[0].attrib['value'], value)

    @patch('be.ldapadmin.orgs_editor.validate_org_info')
    def test_create_org_submit_invalid(self, mock_validator):
        self.request.form = dict(org_info_fixture, id='dk_bridge_club')
        mock_validator.return_value = validation_errors_fixture

        self.ui.create_organisation(self.request)

        mock_validator.assert_called_once_with('dk_bridge_club', org_info_fixture,
                                               create_mode=True)
        self.assertEqual(self.mock_agent.create_org.call_count, 0)
        self.request.RESPONSE.redirect.assert_called_with(
            'URL/create_organisation_html')

        self.request.form = {}
        page = parse_html(self.ui.create_organisation_html(self.request))

        errors = set(["Organisation not created. "
                      "Please correct the errors below."])
        for err_values in validation_errors_fixture.values():
            errors.update(err_values)
        org_info = dict(org_info_fixture, id='dk_bridge_club')
        self._verify_org_form_submit_error(page, org_info, errors)

    def test_edit_org_form(self):
        self.request.form = {'id': 'dk_bridge_club'}
        self.mock_agent.org_info.return_value = dict(org_info_fixture,
                                                     id='dk_bridge_club')

        page = parse_html(self.ui.edit_organisation_html(self.request))

        self.mock_agent.org_info.assert_called_once_with('dk_bridge_club')

        form = page.xpath('//form')[0]
        self.assertEqual(form.attrib['action'], 'URL/edit_organisation')
        self.assertEqual(form.xpath('//input[@name="id"]')[0].attrib['value'],
                         'dk_bridge_club')
        self._verify_org_form_values(form, org_info_fixture)

    @patch('be.ldapadmin.orgs_editor.logged_in_user')
    def test_edit_org_submit(self, logged_user):
        logged_user.return_value = "John Doe"
        self.request.form = dict(org_info_fixture, id='dk_bridge_club')

        self.ui.edit_organisation(self.request)

        self.mock_agent.set_org_info.assert_called_once_with(
            'dk_bridge_club', org_info_fixture)
        self.request.RESPONSE.redirect.assert_called_with(
            'URL/organisation?id=dk_bridge_club')

        self.request.form = {'id': 'dk_bridge_club'}
        page = parse_html(self.ui.organisation(self.request))
        msg = page.xpath('//div[@class="system-msg"]')[0].text
        self.assertTrue(msg.startswith('Organisation saved'))
        logmsg = "John Doe EDITED ORGANISATION dk_bridge_club\n"
        self.assertEqual(self.stream.getvalue(), logmsg)

    @patch('be.ldapadmin.orgs_editor.validate_org_info')
    def test_edit_org_submit_invalid(self, mock_validator):
        self.request.form = dict(org_info_fixture, id='dk_bridge_club')
        self.request.form['id'] = 'dk_bridge_club'
        mock_validator.return_value = validation_errors_fixture

        self.ui.edit_organisation(self.request)

        mock_validator.assert_called_once_with('dk_bridge_club', org_info_fixture)
        self.assertEqual(self.mock_agent.set_org_info.call_count, 0)
        self.request.RESPONSE.redirect.assert_called_with(
            'URL/edit_organisation_html?id=dk_bridge_club')

        self.request.form = {'id': 'dk_bridge_club'}
        page = parse_html(self.ui.edit_organisation_html(self.request))
        self.assertEqual(self.mock_agent.org_info.call_count, 0)

        errors = set(["Organisation not modified. "
                      "Please correct the errors below."])
        for err_values in validation_errors_fixture.values():
            errors.update(err_values)
        org_info = dict(org_info_fixture, id='dk_bridge_club')
        self._verify_org_form_submit_error(page, org_info, errors)

    def test_list_organisations(self):
        self.mock_agent.all_organisations.return_value = {
            'dk_bridge_club': {'name': "Bridge club", 'country': 'dk'},
            'dk_poker_club': {'name': "P\xf8ker club", 'country': 'dk'},
        }

        self.request.get.return_value = None  # no country filter
        page = parse_html(self.ui.index_html(self.request))

        rows = page.xpath('//div[@class="organisation-listing"]//tbody/tr')
        self.assertEqual(len(rows), 2)
        self.assertEqual(plaintext(rows[0].xpath('td')[0]), "Denmark (DK)")
        self.assertEqual(plaintext(rows[0].xpath('td')[1]), "Bridge club")
        self.assertEqual(plaintext(rows[1].xpath('td')[1]), "P\xf8ker club")
        self.assertEqual(rows[0].xpath('td/a')[0].attrib['href'],
                         'URL/organisation?id=dk_bridge_club')

    def test_org_info_page(self):
        self.request.form = {'id': 'dk_bridge_club'}
        self.mock_agent.org_info.return_value = dict(org_info_fixture,
                                                     id='dk_bridge_club')

        page = parse_html(self.ui.organisation(self.request))

        org_h1 = page.xpath('//h1')[0]
        self.assertTrue("Ye olde bridge club" in org_h1.text_content())

        org_table = page.xpath('//table')[0]

        def html_value(label):
            xp = '//tr[td[text()="%s"]]/td[2]' % label
            return org_table.xpath(xp)[0].text

        self.assertEqual(html_value('Name:'), org_info_fixture['name'])
        self.assertEqual(html_value('URL:'), org_info_fixture['url'])
        self.assertEqual(html_value('Phone:'), org_info_fixture['phone'])
        self.assertEqual(html_value('Fax:'), org_info_fixture['fax'])
        self.assertEqual(html_value('Street:'), org_info_fixture['street'])
        self.assertEqual(html_value('PO box:'), org_info_fixture['po_box'])
        self.assertEqual(html_value('Postal code:'),
                         org_info_fixture['postal_code'])
        self.assertEqual(html_value('Locality:'), org_info_fixture['locality'])

    def test_rename_org_page(self):
        self.request.form = {'id': 'dk_bridge_club'}
        self.mock_agent.org_info.return_value = dict(org_info_fixture,
                                                     id='dk_bridge_club')

        page = parse_html(self.ui.rename_organisation_html(self.request))

        form = page.xpath('//form[@name="rename_organisation"]')[0]
        self.assertEqual(form.attrib['action'], 'URL/rename_organisation')
        org_id_input = form.xpath('//input[@name="id"]')[0]
        self.assertEqual(org_id_input.attrib['value'], 'dk_bridge_club')

    @patch('be.ldapadmin.orgs_editor.logged_in_user')
    def test_rename_org_submit(self, logged_user):
        logged_user.return_value = "John Doe"
        self.request.form = {'id': 'dk_bridge_club', 'new_id': 'tunnel_club'}

        self.ui.rename_organisation(self.request)

        self.mock_agent.rename_org.assert_called_once_with('dk_bridge_club',
                                                           'tunnel_club')
        self.request.RESPONSE.redirect.assert_called_with(
            'URL/organisation?id=tunnel_club')

        msg = 'Organisation "dk_bridge_club" renamed to "tunnel_club".'
        self.assertEqual(session_messages(self.request), {'info': [msg]})
        logmsg = "John Doe RENAMED ORGANISATION dk_bridge_club TO tunnel_club\n"
        self.assertEqual(self.stream.getvalue(), logmsg)

    def test_rename_org_submit_fail(self):
        self.request.form = {'id': 'dk_bridge_club', 'new_id': 'tunnel_club'}
        self.mock_agent.rename_org.side_effect = ldapadmin.NameAlreadyExists()

        self.ui.rename_organisation(self.request)

        self.mock_agent.rename_org.assert_called_once_with('dk_bridge_club',
                                                           'tunnel_club')
        self.request.RESPONSE.redirect.assert_called_with(
            'URL/organisation?id=dk_bridge_club')

        msg = ('Organisation "dk_bridge_club" could not be renamed because '
               '"tunnel_club" already exists.')
        self.assertEqual(session_messages(self.request), {'error': [msg]})

    def test_rename_org_submit_crash(self):
        self.request.form = {'id': 'dk_bridge_club', 'new_id': 'tunnel_club'}
        self.mock_agent.rename_org.side_effect = ldapadmin.OrgRenameError()

        self.ui.rename_organisation(self.request)

        self.mock_agent.rename_org.assert_called_once_with('dk_bridge_club',
                                                           'tunnel_club')
        self.request.RESPONSE.redirect.assert_called_with('URL/')

        msg = ('Renaming of "dk_bridge_club" failed mid-way. Some data may be '
               'inconsistent. Please inform a system administrator.')
        self.assertEqual(session_messages(self.request), {'error': [msg]})

    def test_delete_org_page(self):
        import re
        self.request.form = {'id': 'dk_bridge_club'}
        self.mock_agent.org_info.return_value = dict(org_info_fixture,
                                                     id='dk_bridge_club')

        page = parse_html(self.ui.delete_organisation_html(self.request))

        txt = page.xpath('//p[@class="confirm-delete"]')[0].text_content()
        self.assertEqual(re.sub(r'\s+', ' ', txt.strip()),
                         ("Are you sure you want to delete the organisation "
                          "Ye olde bridge club (dk_bridge_club)?"))
        id_input = page.xpath('//form//input[@name="id"]')[0]
        self.assertEqual(id_input.attrib['value'], 'dk_bridge_club')

    @patch('be.ldapadmin.orgs_editor.logged_in_user')
    def test_delete_org_submit(self, logged_user):
        logged_user.return_value = "John Doe"
        self.request.form = {'id': 'dk_bridge_club'}

        self.ui.delete_organisation(self.request)

        self.mock_agent.delete_org.assert_called_once_with('dk_bridge_club')
        self.request.RESPONSE.redirect.assert_called_with('URL/')

        page = parse_html(self.ui.index_html(self.request))
        self.assertEqual(page.xpath('//div[@class="system-msg"]')[0].text,
                         'Organisation "dk_bridge_club" has been deleted.')
        logmsg = "John Doe DELETED ORGANISATION dk_bridge_club\n"
        self.assertEqual(self.stream.getvalue(), logmsg)


class OrganisationsUIMembersTest(unittest.TestCase):
    def setUp(self):
        self.ui = StubbedOrganisationsEditor('organisations')
        self.mock_agent = Mock()
        self.ui._get_ldap_agent = Mock(return_value=self.mock_agent)
        self.request = mock_request()
        user = self.request.AUTHENTICATED_USER
        user.getRoles.return_value = ['Authenticated']
        self.ui.REQUEST = self.request

        user_list = {
            'anne': {'id': 'anne', 'first_name': "Anne",
                     'last_name': "Tester", 'status': ''},
            'jsmith': {'id': 'jsmith', 'first_name': "Joe",
                       'last_name': "Smith", 'status': ''},
        }
        self.mock_agent.members_in_org.return_value = sorted(user_list.keys())
        self.mock_agent.user_info.side_effect = user_list.get
        self.mock_agent.org_info.return_value = dict(org_info_fixture,
                                                     id='dk_bridge_club')

        self.stream = StringIO()
        self.handler = logging.StreamHandler(self.stream)
        self.log = logging.getLogger('orgs_editor')
        self.log.setLevel(logging.INFO)
        for handler in self.log.handlers:
            self.log.removeHandler(handler)
        self.log.addHandler(self.handler)

    def tearDown(self):
        self.log = logging.getLogger('orgs_editor')
        self.log.removeHandler(self.handler)
        self.handler.close()

    def test_enumerate_members(self):
        self.request.form = {'id': 'dk_bridge_club'}

        page = parse_html(self.ui.members_html(self.request))

        self.mock_agent.members_in_org.assert_called_once_with('dk_bridge_club')
        self.mock_agent.user_info.assert_called_with('jsmith')

        form = page.xpath('//form')[0]
        self.assertEqual(form.attrib['action'],
                         'URL/remove_members')
        self.assertEqual(form.xpath('.//input[@name="id"]')[0].attrib['value'],
                         'dk_bridge_club')

        rows = form.xpath('.//table/tbody/tr')
        self.assertTrue("Anne Tester" in rows[0].text_content())
        self.assertTrue("Joe Smith" in rows[1].text_content())

        anne_checkbox = rows[0].xpath('.//input')[0]
        self.assertEqual(anne_checkbox.attrib['name'], 'user_id:list')
        self.assertEqual(anne_checkbox.attrib['value'], 'anne')

    @patch('be.ldapadmin.orgs_editor.logged_in_user')
    def test_remove_members_submit(self, logged_user):
        logged_user.return_value = "John Doe"
        self.request.form = {'id': 'dk_bridge_club', 'user_id': ['jsmith']}

        self.ui.remove_members(self.request)

        self.mock_agent.remove_from_org.assert_called_once_with(
            'dk_bridge_club', ['jsmith'])
        self.request.RESPONSE.redirect.assert_called_with(
            'URL/members_html?id=dk_bridge_club')

        self.request.form = {'id': 'dk_bridge_club'}
        page = parse_html(self.ui.members_html(self.request))
        self.assertEqual(page.xpath('//div[@class="system-msg"]')[0].text,
                         'Removed 1 members from organisation "dk_bridge_club".')
        logmsg = ("John Doe REMOVED MEMBERS ['jsmith'] "
                  "FROM ORGANISATION dk_bridge_club\n")
        self.assertEqual(self.stream.getvalue(), logmsg)

    def test_add_members_html(self):
        self.request.form = {'id': 'dk_bridge_club', 'search_query': u"smith"}
        self.mock_agent.search_user.return_value = [
            {'id': 'anne', 'first_name': "Anne", 'last_name': "Smith"},
            {'id': 'jsmith', 'first_name': "Joe", 'last_name': "Something"},
        ]

        page = parse_html(self.ui.add_members_html(self.request))

        form_search = page.xpath('//form[@name="search-users"]')[0]
        self.assertEqual(form_search.attrib['action'],
                         'URL/add_members_html')
        _xp = './/input[@name="search_query:utf8:string"]'
        self.assertEqual(form_search.xpath(_xp)[0].attrib['value'], u"smith")

        form_add_members = page.xpath('//form[@name="add-members"]')[0]
        self.assertEqual(form_add_members.attrib['action'],
                         'URL/add_members')

        self.mock_agent.search_user.assert_called_once_with(u'smith')
        results_li = form_add_members.xpath('.//ul/li')
        self.assertTrue("Anne Smith" in results_li[0].text_content())
        self.assertTrue("Joe Something" in results_li[1].text_content())

        anne_checkbox = results_li[0].xpath('.//input')[0]
        self.assertEqual(anne_checkbox.attrib['name'], 'user_id:list')
        self.assertEqual(anne_checkbox.attrib['value'], 'anne')

    def test_add_members_submit(self):
        self.request.form = {'id': 'dk_bridge_club', 'user_id': ['jsmith']}

        self.ui.add_members(self.request)

        self.mock_agent.add_to_org.assert_called_once_with(
            'dk_bridge_club', ['jsmith'])
        self.request.RESPONSE.redirect.assert_called_with(
            'URL/members_html?id=dk_bridge_club')

        self.request.form = {'id': 'dk_bridge_club'}
        page = parse_html(self.ui.members_html(self.request))
        self.assertEqual(page.xpath('//div[@class="system-msg"]')[0].text,
                         'Added 1 members to organisation "dk_bridge_club".')


class OrganisationsValidationTest(unittest.TestCase):
    def _test_bad_values(self, name, values, msg):
        for bad_value in values:
            org_info = dict(org_info_fixture, **{name: bad_value})
            err = validate_org_info('myorg', org_info)
            self.assertEqual(err[name], [msg],
                             "Missed bad %s %r" % (name, bad_value))

    def _test_good_values(self, name, values):
        for ok_value in values:
            org_info = dict(org_info_fixture, **{name: ok_value})
            err = validate_org_info('myorg', org_info)
            self.assertTrue(name not in err,
                            "False positive %r %r" % (name, ok_value))

    def _test_phone(self, name, msg):
        bad_values = ['asdf', '1234adsf', '555 1234', '+55 3123 asdf']
        self._test_bad_values(name, bad_values, msg)

        good_values = ['', '+40123456789', '+40 12 34 56 78 9']
        self._test_good_values(name, good_values)

    def test_phone_number(self):
        self._test_phone('phone', VALIDATION_ERRORS['phone'])

    def test_fax_number(self):
        self._test_phone('fax', VALIDATION_ERRORS['fax'])

    def test_postcode(self):
        bad_values = ['123 456', 'DK_1234 33', u'DK 123\xf8456']
        self._test_bad_values('postal_code', bad_values,
                              VALIDATION_ERRORS['postal_code'])
        self._test_good_values('postal_code', ['', 'DK 1Nu 456', 'ro31-23'])

    def test_id(self):
        # the id is checked when creating; it starts with a country code
        for bad_id in ['', '123', 'asdf123', 'my(org)', 'my_\xf8rg', 'a',
                       '_myorg', 'dk_', 'dk-org', 'DK_org']:
            err = validate_org_info(bad_id, dict(org_info_fixture),
                                    create_mode=True)
            self.assertEqual(err['id'], [VALIDATION_ERRORS['id']],
                             "Missed bad org_id %r" % (bad_id,))

        for ok_id in ['dk_org', 'my_org', 'ro_eau_de_web', 'be_a___b']:
            err = validate_org_info(ok_id, dict(org_info_fixture),
                                    create_mode=True)
            self.assertTrue('id' not in err,
                            "False positive org_id %r" % (ok_id,))


XLSX_TYPE = (
    'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


def sheet_rows(data, name):
    wb = load_workbook(BytesIO(data))
    return [list(r) for r in wb[name].iter_rows(values_only=True)]


class OrganisationsExportTest(unittest.TestCase):
    def setUp(self):
        countries = patch.dict(
            'be.ldapadmin.countries.COUNTRIES',
            {'dk': {'code': 'dk', 'name': 'Denmark', 'pub_code': 'DK'}})
        countries.start()
        self.addCleanup(countries.stop)
        self.ui = StubbedOrganisationsEditor('organisations')
        self.agent = Mock()
        self.ui._get_ldap_agent = Mock(return_value=self.agent)
        self.request = mock_request()
        self.request.AUTHENTICATED_USER.getRoles.return_value = [
            'Authenticated']
        self.org = dict(org_info_fixture, id='dk_bridge_club',
                        email='club@example.com')
        users = {
            'anne': {'uid': 'anne', 'first_name': 'Anne',
                     'full_name': 'Anne Tester', 'email': 'anne@example.com'},
            'jsmith': {'uid': 'jsmith', 'first_name': 'Joe',
                       'full_name': 'Joe Smith', 'email': 'joe@example.com'},
        }
        self.agent.all_organisations.return_value = {
            'dk_bridge_club': self.org}
        self.agent.org_info.return_value = self.org
        self.agent.members_in_org.return_value = ['jsmith', 'anne']
        self.agent.user_info.side_effect = users.get

    def headers(self):
        calls = self.request.RESPONSE.setHeader.call_args_list
        return dict(c[0] for c in calls)

    def test_export_organisations(self):
        data = self.ui.export_organisations(self.request)
        headers = self.headers()
        self.assertEqual(headers['Content-Type'], XLSX_TYPE)
        self.assertEqual(headers['Content-Disposition'],
                         'attachment; filename=organisations.xlsx')
        self.assertEqual(headers['Content-Length'], len(data))
        wb = load_workbook(BytesIO(data))
        self.assertEqual(wb.sheetnames, ['Organisations', 'Members'])
        orgs = sheet_rows(data, 'Organisations')
        self.assertEqual(orgs[0], ['Id', 'Name', 'Locality', 'Postal_address',
                                   'Fax', 'Email', 'Members count'])
        self.assertEqual(orgs[2], ['dk_bridge_club', 'Ye olde bridge club',
                                   'K\xf8benhavn', 'Card games road',
                                   '+45 555 9999', 'club@example.com', 2])
        self.assertTrue(wb['Organisations']['A1'].font.bold)
        members = sheet_rows(data, 'Members')
        self.assertEqual(members[0][0], 'Ye olde bridge club')
        self.assertEqual(members[2], ['user id', 'fullname', 'email'])
        self.assertEqual(members[3], ['anne', 'Anne Tester',
                                      'anne@example.com'])

    def test_export_org(self):
        self.request.form = {'id': 'dk_bridge_club'}
        data = self.ui.export_org(self.request)
        headers = self.headers()
        self.assertEqual(headers['Content-Type'], XLSX_TYPE)
        self.assertEqual(headers['Content-Disposition'],
                         'attachment; filename=dk_bridge_club.xlsx')
        wb = load_workbook(BytesIO(data))
        self.assertEqual(wb.sheetnames, ['Organisation Details', 'Users'])
        details = sheet_rows(data, 'Organisation Details')
        self.assertEqual(details[0], ['id', 'name', 'locality',
                                      'postal_address', 'fax', 'email'])
        self.assertEqual(details[2][:2], ['dk_bridge_club',
                                          'Ye olde bridge club'])
        users = sheet_rows(data, 'Users')
        self.assertEqual(users[0], ['user id', 'fullname', 'email'])
        self.assertEqual(users[2], ['anne', 'Anne Tester',
                                    'anne@example.com'])
        self.assertTrue(wb['Users']['A1'].font.bold)
