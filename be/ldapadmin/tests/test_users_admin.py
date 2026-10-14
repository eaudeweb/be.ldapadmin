''' users admin views with a mocked LDAP agent '''
import json
import unittest
from datetime import datetime
from io import BytesIO
from unittest.mock import MagicMock as Mock, call, patch

import xlwt
from lxml.html.soupparser import fromstring
from openpyxl import load_workbook

from be.ldapadmin.db_agent import UsersDB, DISABLE_ACCOUNT
from be.ldapadmin.tests.test_excel_reader import build_xlsx
from be.ldapadmin.tests.test_import_export import read_xlsx
from be.ldapadmin.tests.test_users_ui import StubbedUsersAdmin, user_fixture
from be.ldapadmin.users_admin import (
    BulkUserImporter, ResetUser, MigrateDisabledEmails, SESSION_MESSAGES,
    SESSION_FORM_ERRORS, process_url, generate_password, password_letters,
    _transliterate)


def mock_request(form=None):
    request = Mock()
    request.SESSION = {}
    request.form = dict(form or {})
    request.AUTHENTICATED_USER.getRoles.return_value = ['Authenticated']
    request.AUTHENTICATED_USER.getId.return_value = 'admin'
    return request


def mock_agent():
    agent = Mock()
    # a real context manager: `with agent.new_action:` must fail
    agent.new_action = UsersDB.new_action.__get__(agent)
    agent._update_full_name.side_effect = \
        lambda info: UsersDB._update_full_name(None, info)
    agent.all_organisations.return_value = {
        'dk_org': {'name': 'Org', 'name_native': '', 'country': 'dk'},
        'ro_org': {'name': 'Asociatia', 'name_native': 'Asociația',
                   'country': 'ro'},
    }
    agent.existing_usernames.return_value = iter([])
    agent.existing_emails.return_value = iter([])
    return agent


def messages(request, msg_type):
    return list(request.SESSION.get(SESSION_MESSAGES, {}).get(msg_type, []))


def xls_file(header, rows):
    wb = xlwt.Workbook()
    ws = wb.add_sheet('Sheet 1')
    for col, value in enumerate(header):
        ws.write(0, col, value)
    for row_number, row in enumerate(rows, 1):
        for col, value in enumerate(row):
            ws.write(row_number, col, value)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def xlsx_file(header, rows):
    return build_xlsx([[list(header)] + [list(r) for r in rows]])


XLSX_TYPE = (
    'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


class AdminTestBase(unittest.TestCase):

    def setUp(self):
        self.ui = StubbedUsersAdmin()
        self.agent = mock_agent()
        self.ui._get_ldap_agent = Mock(return_value=self.agent)

    def request(self, form=None):
        request = mock_request(form)
        self.ui.REQUEST = request
        return request


class CreateUserTest(AdminTestBase):

    def test_form_lists_organisations(self):
        request = self.request()
        page = fromstring(self.ui.create_user(request))
        options = page.xpath('//select[@id="edit-organisation"]/option')
        labels = [o.text for o in options]
        self.assertEqual(labels, ['-', 'Asociatia (Asociația, ro_org)',
                                  'Org (dk_org)'])
        self.assertTrue(page.xpath('//input[@name="id:utf8:string"]'))
        password = page.xpath('//input[@name="password:utf8:string"]')
        self.assertEqual(len(password[0].get('value')), 8)

    def test_create(self):
        self.agent.org_exists.return_value = True
        request = self.request({
            'id': 'anne', 'password': 'secret123', 'first_name': 'Anne',
            'last_name': 'M\xfcller', 'email': 'anne@example.com',
            'organisation': 'dk_org', 'submit': 'Create account'})
        self.ui.create_user(request)

        request.RESPONSE.redirect.assert_called_once_with('URL')
        user_id, info = self.agent.create_user.call_args[0]
        self.assertEqual(user_id, 'anne')
        self.assertEqual(info['first_name'], 'Anne')
        self.assertEqual(info['email'], 'anne@example.com')
        self.assertIn('Mueller', info['search_helper'].split())
        self.assertNotIn('organisation', info)
        self.agent.set_user_password.assert_called_once_with(
            'anne', None, 'secret123')
        self.agent.add_to_org.assert_called_once_with('dk_org', ['anne'])
        self.assertTrue(messages(request, 'info')[0].startswith(
            'User anne created'))

    def test_create_invalid(self):
        request = self.request({
            'id': 'Anne!', 'first_name': 'Anne', 'last_name': 'Tester',
            'email': 'not-an-email', 'submit': 'Create account'})
        page = fromstring(self.ui.create_user(request))
        self.assertFalse(self.agent.create_user.called)
        errors = {e.get('id'): e.text
                  for e in page.xpath('//p[@class="error-hint"]')}
        self.assertIn('error-edit-id', errors)
        self.assertEqual(errors['error-edit-email'],
                         'Invalid email format not-an-email')
        self.assertIn('Please correct the errors below and try again.',
                      page.text_content())


class EditUserActionTest(AdminTestBase):

    def test_save(self):
        self.agent.set_user_info.return_value = True
        request = self.request({
            'id': 'anne', 'first_name': 'Anne', 'last_name': 'Tester',
            'email': 'anne@example.com, anne@example.org',
            'organisation': 'dk_org'})
        self.ui.edit_user_action(request)
        user_id, info = self.agent.set_user_info.call_args[0]
        self.assertEqual(user_id, 'anne')
        self.assertEqual(info['organisation'], ['dk_org'])
        self.assertEqual(info['email'],
                         ['anne@example.com', 'anne@example.org'])
        self.assertTrue(messages(request, 'info')[0].startswith(
            'Profile saved'))
        request.RESPONSE.redirect.assert_called_once_with(
            'URL/edit_user?id=anne')

    def test_invalid(self):
        request = self.request({
            'id': 'anne', 'first_name': 'Anne', 'last_name': 'Tester',
            'email': 'bad'})
        self.ui.edit_user_action(request)
        self.assertFalse(self.agent.set_user_info.called)
        self.assertEqual(request.SESSION[SESSION_FORM_ERRORS],
                         {'email': 'Invalid email format bad'})


class ConfirmationPagesTest(AdminTestBase):

    def setUp(self):
        super().setUp()
        self.agent.user_info.return_value = user_fixture()

    def check_page(self, method, heading):
        request = self.request({'id': 'anne'})
        page = fromstring(getattr(self.ui, method)(request))
        self.assertIn(heading, ' '.join(page.xpath('//h1')[0].itertext()))
        self.assertIn('Anne Tester', page.text_content())
        self.assertEqual(page.xpath('//input[@name="id"]')[0].get('value'),
                         'anne')

    def test_delete_page(self):
        self.check_page('delete_user', 'Delete user')

    def test_disable_page(self):
        self.check_page('disable_user', 'Disable user')

    def test_enable_page(self):
        self.check_page('enable_user', 'Enable user')

    def test_enable_page_offers_restore_choice(self):
        page = fromstring(self.ui.enable_user(self.request({'id': 'anne'})))
        radios = page.xpath('//input[@name="restore_roles"]')
        self.assertEqual([(r.get('type'), r.get('value')) for r in radios],
                         [('radio', 'on'), ('radio', '')])
        self.assertEqual(radios[0].get('checked'), 'checked')
        self.assertIn('Restore all', page.text_content())
        self.assertIn('Restore nothing', page.text_content())

    def test_delete_action(self):
        request = self.request({'id': 'anne'})
        self.ui.delete_user_action(request)
        self.agent.delete_user.assert_called_once_with('anne')
        self.assertEqual(messages(request, 'info'),
                         ['User "anne" has been deleted.'])

    def test_disable_action(self):
        request = self.request({'id': 'anne'})
        self.ui.disable_user_action(request)
        self.agent.disable_user.assert_called_once_with('anne')
        self.assertIn('has been disabled', messages(request, 'info')[0])

    @patch('be.ldapadmin.users_admin.getUtility')
    def test_enable_action_sends_email(self, getUtility):
        self.ui.unrestrictedTraverse = Mock(
            return_value=Mock(absolute_url=Mock(return_value='SITE')))
        request = self.request({'id': 'anne', 'restore_roles': 'on'})
        self.agent.enable_user.return_value = {
            'organisations': ['dk_org'], 'roles': ['places', 'places-bank'],
            'roles_permittedPerson': ['places'], 'roles_owner': [],
            'skipped_roles': ['gone']}
        self.ui.enable_user_action(request)
        self.agent.enable_user.assert_called_once_with(
            'anne', restore_roles='on')
        addr_from, addr_to, body = getUtility.return_value.send.call_args[0]
        self.assertEqual(addr_to, ['anne@example.com'])
        self.assertIn('Your username is anne.', body)
        self.assertIn('SITE/password-reset', body)
        message = messages(request, 'info')[0]
        self.assertIn('restored: organisations dk_org; roles places, '
                      'places-bank; permittedPerson in roles places; '
                      'owner of roles -; skipped deleted roles gone', message)

    @patch('be.ldapadmin.users_admin.getUtility')
    def test_enable_action_without_restore(self, getUtility):
        self.ui.unrestrictedTraverse = Mock(
            return_value=Mock(absolute_url=Mock(return_value='SITE')))
        request = self.request({'id': 'anne', 'restore_roles': ''})
        self.ui.enable_user_action(request)
        self.agent.enable_user.assert_called_once_with(
            'anne', restore_roles='')
        self.assertIn('previous roles NOT restored',
                      messages(request, 'info')[0])


class FindDuplicatesTest(AdminTestBase):

    @patch('be.ldapadmin.users_admin.get_duplicates_by_name')
    def test_duplicates(self, by_name):
        by_name.return_value = ['uid=bob,ou=Users', 'uid=anne,ou=Users']
        self.agent._user_id.side_effect = lambda dn: dn[4:].split(',')[0]
        self.agent.search_user_by_email.return_value = [user_fixture()]
        self.agent.search_users_by_uid.return_value = [user_fixture(
            id='bob', uid='bob', full_name='Bob Builder')]
        request = self.request({'first_name': 'Anne', 'last_name': 'Tester',
                                'email': 'anne@example.com'})
        page = fromstring(self.ui.find_duplicates(request))
        by_name.assert_called_once_with('Anne Tester')
        self.assertEqual(self.agent.search_users_by_uid.call_args[0][0],
                         {'bob'})
        text = page.text_content()
        self.assertIn('Anne Tester', text)
        self.assertIn('Bob Builder', text)

    def test_missing_fields(self):
        request = self.request({'first_name': 'Anne'})
        out = self.ui.find_duplicates(request)
        self.assertNotIn('Suggested similar users', out)
        self.assertFalse(self.agent.search_user_by_email.called)


class BulkChecksTest(AdminTestBase):

    def test_bulk_check_username(self):
        self.agent.existing_usernames.return_value = iter(['anne'])
        request = self.request({'usernames': 'anne, Bad-Name\nbob'})
        page = fromstring(self.ui.bulk_check_username(request))
        self.assertEqual(page.xpath('//textarea[@id="id_valid"]')[0].text,
                         'bob')
        taken = [t.text for t in page.xpath('//textarea[@id="id_taken"]')]
        self.assertEqual(taken, ['anne', 'Bad-Name'])
        self.assertEqual(sorted(self.agent.existing_usernames.call_args[0][0]),
                         ['anne', 'bob'])

    def test_bulk_check_email(self):
        self.agent.existing_emails.return_value = iter(['anne@example.com'])
        request = self.request({'emails': (
            'Anne@example.com bob@example.com bob@EXAMPLE.com wrong')})
        page = fromstring(self.ui.bulk_check_email(request))
        self.assertEqual(page.xpath('//textarea[@id="id_valid"]')[0].text,
                         'bob@example.com')
        taken = [t.text for t in page.xpath('//textarea[@id="id_taken"]')]
        self.assertEqual(taken, ['anne@example.com', 'wrong',
                                 'bob@EXAMPLE.com'])

    def test_bulk_get_emails(self):
        members = {'dk_org': ['anne'], 'ro_org': ['anne', 'bob']}
        self.agent.members_in_org.side_effect = members.get
        self.agent.user_info.side_effect = lambda uid: user_fixture(
            id=uid, email='%s@example.com' % uid)
        emails = json.loads(self.ui.bulk_get_emails(self.request()))
        self.assertEqual(sorted(emails),
                         ['anne@example.com', 'bob@example.com'])


class StatisticsTest(AdminTestBase):

    def test_statistics(self):
        self.agent.conn.allresults.return_value = [
            (None, [('uid=anne', {}), ('uid=bob', {})], 1, []),
            (None, [('uid=carl', {})], 1, []),
        ]
        years = {'uid=anne': 2010, 'uid=bob': 2012, 'uid=carl': 2010}
        self.agent._unpack_user_info.side_effect = lambda dn, attr: {
            'createTimestamp': datetime(years[dn], 1, 1)}
        page = fromstring(self.ui.get_statistics(self.request()))
        cells = [[td.text for td in tr.xpath('td')]
                 for tr in page.xpath('//tbody/tr')]
        self.assertEqual(cells, [['2010', '2'], ['2012', '1']])


class OrgsInCountryTest(AdminTestBase):

    @patch('be.ldapadmin.countries.COUNTRIES', {
        'dk': {'code': 'dk', 'name': 'Denmark'},
        'ro': {'code': 'ro', 'name': 'Romania'}})
    def test_orgs_in_country(self):
        self.request()
        self.assertEqual(list(self.ui.orgs_in_country('ro')), ['ro_org'])


class EmailsTest(AdminTestBase):

    def setUp(self):
        super().setUp()
        self.ui.unrestrictedTraverse = Mock(
            return_value=Mock(absolute_url=Mock(return_value='SITE')))
        self.request()

    @patch('be.ldapadmin.users_admin._send_email')
    def test_confirmation_email(self, send):
        self.ui.send_confirmation_email({
            'id': 'anne', 'first_name': 'Anne', 'email': 'anne@example.com'})
        addr_from, addr_to, message = send.call_args[0]
        self.assertEqual(addr_to, 'anne@example.com')
        self.assertIn('`anne`', message['Subject'])
        body = message.get_payload(decode=True).decode('utf-8')
        self.assertIn('Hello Anne,', body)
        self.assertIn('user ID `anne`', body)
        self.assertIn('SITE/password-reset', body)

    @patch('be.ldapadmin.users_admin._send_email')
    def test_create_user_sends_helpdesk_email(self, send):
        info = {'id': 'anne', 'password': 'pw', 'first_name': 'Anne',
                'last_name': 'Tester', 'full_name': 'Anne Tester',
                'organisation': 'dk_org'}
        sent = []
        self.agent.create_user.side_effect = lambda uid, i: sent.append(
            dict(i))
        user_id = self.ui._create_user(self.agent, info,
                                       send_helpdesk_email=True)
        self.assertEqual(user_id, 'anne')
        self.assertEqual(len(sent), 1)
        self.assertNotIn('password', sent[0])
        self.assertNotIn('id', sent[0])
        self.assertEqual(info['password'], 'pw')
        message = send.call_args[0][2]
        body = message.get_payload(decode=True).decode('utf-8')
        self.assertIn('created by admin', body)
        self.assertIn('User id: anne', body)
        self.assertIn('Member of organisation: dk_org', body)


class BulkImportTest(AdminTestBase):

    def importer(self, form=None):
        request = self.request(form)
        return BulkUserImporter(self.ui, request), request

    def test_download_template(self):
        view, request = self.importer({'download_template': '1'})
        data = view()
        rows, merged = read_xlsx(data)
        self.assertEqual(rows[0][1], BulkUserImporter.TEMPLATE_COLUMNS)
        headers = dict(c[0] for c in request.RESPONSE.setHeader.call_args_list)
        self.assertEqual(headers['Content-Length'], len(data))
        self.assertEqual(headers['Content-Type'], XLSX_TYPE)
        self.assertEqual(headers['Content-Disposition'],
                         'attachment; filename=create_users_template.xlsx')
        sheet = load_workbook(BytesIO(data)).worksheets[0]
        header = [c.value for c in sheet[1]]
        self.assertEqual(header, BulkUserImporter.TEMPLATE_COLUMNS)

    def template_rows(self, *users):
        rows = []
        for first, last, email in users:
            row = [''] * len(BulkUserImporter.TEMPLATE_COLUMNS)
            row[0], row[1], row[4] = first, last, email
            row[6] = 'example.com'
            row[8] = '+45 33 12 34 56'
            row[11] = 'dk_org'
            row[13] = 'testing'
            rows.append(row)
        return rows

    def test_read_xlsx(self):
        view, request = self.importer()
        data = xlsx_file(BulkUserImporter.TEMPLATE_COLUMNS, self.template_rows(
            ('Anne', 'Tester', ' Anne@Example.com '),
            ('Anna', 'Tester', 'anna@example.com')))
        rows = view.read_xls(data)
        self.assertEqual([r['id'] for r in rows], ['testeann', 'testeana'])
        self.assertEqual(rows[0]['email'], 'anne@example.com')
        self.assertEqual(rows[0]['url'], 'http://example.com')
        self.assertEqual(rows[0]['organisation'], 'dk_org')

    def test_read_xlsx_wrong_columns(self):
        view, request = self.importer()
        with self.assertRaises(ValueError):
            view.read_xls(xlsx_file(['First Name*'], [['Anne']]))

    def test_read_xlsx_number_cell(self):
        view, request = self.importer()
        rows = self.template_rows(('Anne', 'Tester', 'anne@example.com'))
        rows[0][8] = 4533123456
        with self.assertRaises(ValueError) as cm:
            view.read_xls(xlsx_file(BulkUserImporter.TEMPLATE_COLUMNS, rows))
        self.assertIn('format all cells as text', str(cm.exception))

    def test_bulk_create_xlsx(self):
        self.ui.send_confirmation_email = Mock()
        self.ui.send_password_reset_email = Mock()
        self.agent.org_exists.return_value = True
        data = xlsx_file(BulkUserImporter.TEMPLATE_COLUMNS, self.template_rows(
            ('Anne', 'Tester', 'anne@example.com')))
        view, request = self.importer({'bulk_create': '1',
                                       'data': BytesIO(data)})
        with patch('be.ldapadmin.users_admin._send_email'):
            text = fromstring(view()).text_content()
        created = [c[0][0] for c in self.agent.create_user.call_args_list]
        self.assertEqual(created, ['testeann'], text)

    @patch('be.ldapadmin.users_admin._send_email')
    def test_bulk_create(self, send):
        self.ui.send_confirmation_email = Mock()
        self.ui.send_password_reset_email = Mock()
        self.agent.org_exists.return_value = True
        self.agent.existing_emails.return_value = iter(['bob@example.com'])
        data = xlsx_file(BulkUserImporter.TEMPLATE_COLUMNS, self.template_rows(
            ('Anne', 'Tester', 'anne@example.com'),
            ('Bob', 'Builder', 'bob@example.com'),
            ('Carl', 'Bad', 'carl-at-example')))
        view, request = self.importer({'bulk_create': '1',
                                       'data': BytesIO(data)})
        text = fromstring(view()).text_content()
        created = [c[0][0] for c in self.agent.create_user.call_args_list]
        self.assertEqual(created, ['testeann'], text)
        self.assertEqual(self.agent.add_to_org.call_args_list,
                         [call('dk_org', ['testeann'])], text)
        self.assertIn('User(s) Anne Tester (testeann) successfully created.',
                      text)
        self.assertIn('email at row 3: Invalid email format carl-at-example',
                      text)
        self.assertIn('The following email is already in database: '
                      'bob@example.com', text)
        self.ui.send_confirmation_email.assert_called_once()
        self.assertEqual(send.call_count, 1)

    def test_bulk_create_xls_rejected(self):
        data = xls_file(BulkUserImporter.TEMPLATE_COLUMNS, self.template_rows(
            ('Anne', 'Tester', 'anne@example.com')))
        view, request = self.importer({'bulk_create': '1',
                                       'data': BytesIO(data)})
        self.assertIn('Invalid Excel file', view())
        self.assertFalse(self.agent.create_user.called)

    def test_bulk_create_invalid_file(self):
        view, request = self.importer({'bulk_create': '1',
                                       'data': BytesIO(b'not excel')})
        self.assertIn('Invalid Excel file', view())
        self.assertFalse(self.agent.create_user.called)


class OtherViewsTest(AdminTestBase):

    def test_reset_user(self):
        request = self.request({'id': 'anne', 'submit': '1'})
        ResetUser(self.ui, request)()
        self.agent.reset_user_roles.assert_called_once_with('anne')
        request.RESPONSE.redirect.assert_called_once_with(
            'URL/edit_user?id=anne')

    def test_migrate_disabled_emails(self):
        metadata = [{'action': DISABLE_ACCOUNT,
                     'data': {'email': 'anne@example.com'}}]
        self.agent.get_disabled_users.return_value = [
            {'id': 'anne', 'metadata': json.dumps(metadata), 'email': ''}]
        self.agent._get_email_for_disabled_user.side_effect = \
            lambda meta: UsersDB._get_email_for_disabled_user(None, meta)
        request = self.request()
        self.assertEqual(MigrateDisabledEmails(self.ui, request)(), 'done')
        user_id, info = self.agent.set_user_info.call_args[0]
        self.assertEqual(info['email'], 'anne@example.com')


class HelpersTest(unittest.TestCase):

    def test_process_url(self):
        self.assertEqual(process_url('example.com'), 'http://example.com')
        self.assertEqual(process_url('https://x.eu'), 'https://x.eu')
        self.assertEqual(process_url(''), '')

    def test_generate_password(self):
        password = generate_password()
        self.assertEqual(len(password), 8)
        self.assertTrue(set(password) <= set(password_letters))

    def test_transliterate(self):
        helper = _transliterate('J\xfcrgen', 'Stra\xdfe', '', '')
        words = helper.split()
        self.assertIn('Jurgen', words)
        self.assertIn('Strasse', words)
        self.assertEqual(words, sorted(set(words)))
